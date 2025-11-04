"""Excel report generator using openpyxl."""

from pathlib import Path
from typing import Any, Union

from .report_generator import BaseReportGenerator  # Import base class

try:
    import openpyxl

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False
    openpyxl = None  # type: ignore


class ExcelReportGenerator(BaseReportGenerator):  # Inheritance
    """Concrete Excel implementation of the report generator interface."""

    report_type = "excel"
    file_extension = ".xlsx"

    def _render(self, data: dict[str, Any]) -> Path:  # Polymorphism
        """Create an XLSX validation report from the supplied data."""
        if not HAS_OPENPYXL or openpyxl is None:
            msg = "openpyxl is required for Excel report generation"
            raise ImportError(msg)

        excel_file = self.output_dir / "validation_report.xlsx"
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
        except Exception as e:
            raise RuntimeError(f"Cannot create Excel workbook: {e}") from e
        try:
            ws.title = "Validation"  # type: ignore
            # Headers
            ws["A1"] = "Metric"  # type: ignore
            ws["B1"] = "Value"  # type: ignore

            # Data
            status = self._determine_validation_status(data)
            metrics: list[tuple[str, Union[int, str]]] = [
                ("Pages", data.get("pages", 0)),
                ("Content Items", data.get("content_items", 0)),
                ("TOC Entries", data.get("toc_entries", 0)),
                ("Status", status),
            ]

            for i, (metric, value) in enumerate(metrics, 2):
                ws[f"A{i}"] = metric  # type: ignore
                ws[f"B{i}"] = value  # type: ignore
        except Exception as e:
            raise RuntimeError(f"Cannot write Excel data: {e}") from e

        try:
            wb.save(excel_file)
        except OSError as e:
            raise RuntimeError(f"Cannot save Excel file: {e}") from e
        return excel_file
